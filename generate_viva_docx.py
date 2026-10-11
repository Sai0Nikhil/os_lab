import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Load markdown file and parse Q&A
qa_pairs = []
with open("/home/sai-nikhil/os_lab/VIVA_QUESTIONS_ANSWERS.md", "r") as f:
    lines = f.readlines()

current_q = None
current_a = []

for line in lines:
    line_s = line.strip()
    if line_s.startswith("### "):
        if current_q:
            qa_pairs.append((current_q, " ".join(current_a).strip()))
        current_q = line_s[4:]
        current_a = []
    elif current_q and line_s and not line_s.startswith("---") and not line_s.startswith("# "):
        current_a.append(line_s)

if current_q:
    qa_pairs.append((current_q, " ".join(current_a).strip()))

print(f"Total parsed Q&A pairs: {len(qa_pairs)}")

doc = docx.Document()
for s in doc.sections:
    s.top_margin = Inches(0.8)
    s.bottom_margin = Inches(0.8)
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

PRIMARY = RGBColor(24, 76, 120)
ACCENT = RGBColor(41, 128, 185)

# Title
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = t.add_run("OPERATING SYSTEMS LABORATORY")
tr.bold = True
tr.font.size = Pt(20)
tr.font.name = "Arial"
tr.font.color.rgb = PRIMARY

st = doc.add_paragraph()
st.alignment = WD_ALIGN_PARAGRAPH.CENTER
strun = st.add_run("Viva Voce Questions and Comprehensive Answers (Q1 – Q60)\n")
strun.font.size = Pt(13)
strun.font.name = "Arial"
strun.font.italic = True
strun.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph("─" * 70)

for idx, (q, a) in enumerate(qa_pairs, 1):
    qp = doc.add_paragraph()
    qp.paragraph_format.space_before = Pt(8)
    qp.paragraph_format.space_after = Pt(2)
    
    q_num_run = qp.add_run(f"Question: ")
    q_num_run.bold = True
    q_num_run.font.name = "Arial"
    q_num_run.font.size = Pt(11)
    q_num_run.font.color.rgb = PRIMARY
    
    q_text_run = qp.add_run(q)
    q_text_run.bold = True
    q_text_run.font.name = "Arial"
    q_text_run.font.size = Pt(11)
    q_text_run.font.color.rgb = RGBColor(20, 20, 20)
    
    ap = doc.add_paragraph()
    ap.paragraph_format.space_before = Pt(2)
    ap.paragraph_format.space_after = Pt(8)
    ap.paragraph_format.left_indent = Inches(0.2)
    
    ans_lbl = ap.add_run("Answer: ")
    ans_lbl.bold = True
    ans_lbl.font.name = "Arial"
    ans_lbl.font.size = Pt(10.5)
    ans_lbl.font.color.rgb = ACCENT
    
    # Strip markdown formatting like ** or $$
    clean_a = a.replace("**", "").replace("$$", "").replace("\\text", "").replace("{", "").replace("}", "")
    ans_run = ap.add_run(clean_a)
    ans_run.font.name = "Arial"
    ans_run.font.size = Pt(10.5)
    ans_run.font.color.rgb = RGBColor(40, 40, 40)

out_path = "/home/sai-nikhil/os_lab/OS_Lab_Viva_Questions_Answers.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
